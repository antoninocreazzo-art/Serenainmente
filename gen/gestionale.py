#!/usr/bin/env python3
"""Genera scarica/gestionale-di-famiglia.xlsx — il gestionale di famiglia di Serena, versione 2.0.
Uso: python3 gen/gestionale.py  (viene chiamato anche dal generatore del sito)

Regole di compatibilità (Excel 2010+, Excel per telefono, Google Fogli, LibreOffice):
niente FILTER/SORT/XLOOKUP/LET, niente macro. Gli elenchi ordinati si fanno con
SMALL + INDEX/MATCH su una scheda nascosta «Calcoli» che impila Scadenze, Visite e Documenti."""
import os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.environ.get('GEST_OUT') or os.path.join(ROOT, 'scarica', 'gestionale-di-famiglia.xlsx')
VERSIONE = '2.0'

MALVA, INK, MUTED, LINE, BLU = '97606F', '3A2C30', '7A6569', 'EADCDF', '0000FF'
ROSA, SALVIA, SALVIA_T, SABBIA, SABBIA_T, CREMA = 'F4E9EC', 'E8EDE4', '4D5A48', 'EFE6DC', '6B5644', 'FBF6F4'
F = 'Arial'
thin = Side(style='thin', color=LINE)
H = Font(name=F, bold=True, color='FFFFFF', size=10)
IN = Font(name=F, color=BLU, size=10)          # celle da compilare
FX = Font(name=F, color=INK, size=10)          # formule
TITLE = Font(name='Georgia', bold=True, size=18, color=INK)
SUB = Font(name=F, italic=True, size=10, color=MUTED)
EX = Font(name=F, italic=True, size=9, color=MUTED)
fill = lambda c: PatternFill('solid', fgColor=c)
FILL_H, FILL_ALT = fill(MALVA), fill(CREMA)

N = 300          # righe pronte per Scadenze, Visite, Documenti
NF = 30          # persone in Famiglia
R0 = 5           # prima riga dati in ogni scheda
RN = R0 + N - 1  # ultima riga dati

wb = Workbook()
oggi = datetime.date.today()
d = lambda n: oggi + datetime.timedelta(days=n)

def sheet(title, headline, sub, cols):
    """cols: lista di (intestazione, larghezza, 'in'|'fx'|'hide')."""
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws['A1'] = headline; ws['A1'].font = TITLE; ws.row_dimensions[1].height = 28
    ws['A2'] = sub; ws['A2'].font = SUB
    for i, (h, w, kind) in enumerate(cols, start=1):
        c = ws.cell(row=R0 - 1, column=i, value=h); c.font = H; c.fill = FILL_H
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[L(i)].width = w
        if kind == 'hide': ws.column_dimensions[L(i)].hidden = True
    ws.row_dimensions[R0 - 1].height = 32
    ws.freeze_panes = f'B{R0}'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

def body(ws, cols, n, height=20):
    for r in range(R0, R0 + n):
        for i, (_, _, kind) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if kind == 'in':
                c.font = IN
                if (r - R0) % 2: c.fill = FILL_ALT
            else:
                c.font = FX
        ws.row_dimensions[r].height = height

def fmt(ws, col, n, f='DD/MM/YYYY'):
    for r in range(R0, R0 + n): ws[f'{col}{r}'].number_format = f

def lista(ws, col, n, options=None, ref=None, strict=True):
    f1 = ref if ref else '"' + ','.join(options) + '"'
    dv = DataValidation(type='list', formula1=f1, allow_blank=True, showErrorMessage=strict)
    if not strict:
        dv.promptTitle = 'Scegli o scrivi'; dv.prompt = 'Scegli dall\'elenco della scheda Famiglia, oppure scrivi (es. «Tutti»).'
    ws.add_data_validation(dv); dv.add(f'{col}{R0}:{col}{R0 + n - 1}')

PERSONE = f"Famiglia!$A${R0}:$A${R0 + NF - 1}"

def esempio(ws, r, values):
    for i, v in enumerate(values, start=1):
        if v is not None: ws.cell(row=r, column=i, value=v)

# ======================= Famiglia =======================
cols = [('Nome', 16, 'in'), ('Ruolo', 13, 'in'), ('Data di nascita', 14, 'in'), ('Età', 10, 'fx'),
        ('Pediatra o medico di famiglia', 26, 'in'), ('Allergie note (facoltativo)', 22, 'in'),
        ('Gruppo sanguigno (facoltativo)', 14, 'in'), ('Fascicolo sanitario', 22, 'in'), ('Note', 28, 'in')]
ws = sheet('Famiglia', 'La tua famiglia', 'Scrivi qui i nomi una volta sola: compaiono nei menu di tutte le altre schede. L\'età si calcola da sola.', cols)
body(ws, cols, NF); fmt(ws, 'C', NF)
lista(ws, 'B', NF, ['Mamma', 'Papà', 'Figlio/a', 'Nonno/a', 'Altra persona di cui mi occupo', 'Altro'])
lista(ws, 'G', NF, ['Non so', '0+', '0-', 'A+', 'A-', 'B+', 'B-', 'AB+', 'AB-'])
lista(ws, 'H', NF, ['Attivo con consenso', 'Attivo senza consenso', 'Da attivare (figlio)', 'Non so'])
for r in range(R0, R0 + NF):
    ws[f'D{r}'] = f'=IF(C{r}="","",IF(DATEDIF(C{r},TODAY(),"y")=0,DATEDIF(C{r},TODAY(),"m")&" mesi",DATEDIF(C{r},TODAY(),"y")&" anni"))'
esempio(ws, R0, ['Sofia', 'Figlio/a', datetime.date(oggi.year - 2, 3, 14), None, 'Dott.ssa Rossi · 010 000000', 'Nessuna nota', 'Non so', 'Da attivare (figlio)', 'Esempio: cancella e scrivi i tuoi dati'])
esempio(ws, R0 + 1, ['Luca', 'Papà', datetime.date(oggi.year - 36, 9, 2), None, 'Dott. Bianchi', None, None, 'Attivo con consenso', None])
ws.cell(row=R0, column=9).font = EX
ws[f'A{R0 + NF + 1}'] = 'Scrivi solo quello che ti serve davvero: questo file non è un archivio medico protetto. Per terapie e diagnosi il posto giusto è il Fascicolo sanitario.'
ws[f'A{R0 + NF + 1}'].font = EX

# ======================= Scadenze =======================
cols = [('Cosa', 32, 'in'), ('Per chi', 14, 'in'), ('Data', 13, 'in'), ('Si ripete', 13, 'in'), ('Prossima volta', 14, 'fx'),
        ('Giorni', 8, 'fx'), ('Stato', 16, 'fx'), ('Fatto', 8, 'in'), ('Chi se ne occupa', 16, 'in'), ('Priorità', 11, 'in'),
        ('Note', 30, 'in'), ('mesi', 6, 'hide')]
ws = sheet('Scadenze', 'Scadenze e cose da fare', 'Richiami, rinnovi, iscrizioni, pagamenti. Se si ripete, la «Prossima volta» si sposta da sola. Spunta «Fatto» quando è chiusa.', cols)
body(ws, cols, N); fmt(ws, 'C', N); fmt(ws, 'E', N)
lista(ws, 'B', N, ref=PERSONE, strict=False); lista(ws, 'I', N, ref=PERSONE, strict=False)
lista(ws, 'D', N, ['Mai', 'Ogni mese', 'Ogni 3 mesi', 'Ogni 6 mesi', 'Ogni anno'])
lista(ws, 'H', N, ['☐', '☑']); lista(ws, 'J', N, ['Alta', 'Normale'])
MESI = "((YEAR(TODAY())-YEAR(C{r}))*12+MONTH(TODAY())-MONTH(C{r}))"
for r in range(R0, RN + 1):
    m = MESI.format(r=r)
    ws[f'L{r}'] = f'=IF(D{r}="Ogni mese",1,IF(D{r}="Ogni 3 mesi",3,IF(D{r}="Ogni 6 mesi",6,IF(D{r}="Ogni anno",12,0))))'
    ws[f'E{r}'] = (f'=IF(C{r}="","",IF(OR(L{r}=0,C{r}>=TODAY()),C{r},'
                   f'IF(EDATE(C{r},INT({m}/L{r})*L{r})>=TODAY(),EDATE(C{r},INT({m}/L{r})*L{r}),EDATE(C{r},(INT({m}/L{r})+1)*L{r}))))')
    ws[f'F{r}'] = f'=IF(OR(E{r}="",H{r}="☑"),"",E{r}-TODAY())'
    ws[f'G{r}'] = f'=IF(H{r}="☑","Fatto",IF(E{r}="","",IF(F{r}<0,"Scaduta",IF(F{r}<=7,"Questa settimana",IF(F{r}<=30,"Entro 30 giorni","Più avanti")))))'
    ws[f'H{r}'].alignment = Alignment(horizontal='center', vertical='center')
rng = f'G{R0}:G{RN}'
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'G{R0}="Scaduta"'], fill=fill('F8D7DA'), font=Font(name=F, color='8A2430', bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'G{R0}="Questa settimana"'], fill=fill('FFF0C2'), font=Font(name=F, color='7A5A00', bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'G{R0}="Fatto"'], font=Font(name=F, color='9A8A8E')))
ws.conditional_formatting.add(f'A{R0}:K{RN}', FormulaRule(formula=[f'$H{R0}="☑"'], font=Font(name=F, color='9A8A8E', strike=True)))
esempio(ws, R0,     ['Richiamo vaccino (chiedere al pediatra quale)', 'Sofia', d(21), 'Mai', None, None, None, '☐', 'Sofia', 'Alta', 'Esempio: cancella e scrivi i tuoi dati'])
esempio(ws, R0 + 1, ['Retta del nido', 'Sofia', d(-40), 'Ogni mese', None, None, None, '☐', 'Luca', 'Normale', 'Si ripete: la data si sposta da sola'])
esempio(ws, R0 + 2, ['Iscrizione al centro estivo', 'Sofia', d(-10), 'Mai', None, None, None, '☑', 'Luca', 'Normale', 'Fatto: esce dai promemoria'])
ws[f'I{R0}'] = 'Mamma'
for r in (R0, R0 + 1, R0 + 2): ws[f'K{r}'].font = EX

# ======================= Visite =======================
cols = [('Data', 12, 'in'), ('Ora', 8, 'in'), ('Per chi', 13, 'in'), ('Medico o struttura', 24, 'in'), ('Stato', 14, 'in'),
        ('Disdire entro', 13, 'fx'), ('Disdire entro: data diversa (se la struttura dice altro)', 18, 'in'), ('Ticket', 12, 'in'),
        ('Priorità ricetta', 13, 'in'), ('Scadenza ricetta', 13, 'in'), ('N. prenotazione', 16, 'in'),
        ('Domanda 1 (la più importante)', 28, 'in'), ('Domanda 2', 24, 'in'), ('Domanda 3', 24, 'in'), ('Cosa ho notato', 26, 'in'),
        ('Risposte', 30, 'in'), ('Cosa facciamo', 26, 'in'), ('Richiamare se o quando', 22, 'in')]
ws = sheet('Visite', 'Visite: prima, durante e dopo', 'Le prime colonne servono ogni giorno; le domande più a destra. «Disdire entro» propone due giorni lavorativi prima: se la tua struttura dice diverso, scrivi la data nella colonna accanto.', cols)
body(ws, cols, N, height=34); fmt(ws, 'A', N); fmt(ws, 'F', N); fmt(ws, 'G', N); fmt(ws, 'J', N); fmt(ws, 'B', N, 'HH:MM')
lista(ws, 'C', N, ref=PERSONE, strict=False)
lista(ws, 'E', N, ['Programmata', 'Completata', 'Annullata'])
lista(ws, 'H', N, ['Da pagare', 'Pagato', 'Esente', 'Privata'])
lista(ws, 'I', N, ['U (72 ore)', 'B (10 giorni)', 'D (30/60 giorni)', 'P (120 giorni)', 'Controllo', 'Privata'])
for r in range(R0, RN + 1):
    ws[f'F{r}'] = f'=IF(A{r}="","",IF(G{r}<>"",G{r},WORKDAY(A{r},-2)))'
chiusa = f'OR($E{R0}="Completata",$E{R0}="Annullata")'
ws.conditional_formatting.add(f'A{R0}:R{RN}', FormulaRule(formula=[f'$E{R0}="Annullata"'], font=Font(name=F, color='9A8A8E', strike=True)))
ws.conditional_formatting.add(f'A{R0}:R{RN}', FormulaRule(formula=[f'$E{R0}="Completata"'], font=Font(name=F, color='9A8A8E')))
ws.conditional_formatting.add(f'F{R0}:F{RN}', FormulaRule(formula=[f'AND(ISNUMBER(F{R0}),NOT({chiusa}),F{R0}>=TODAY(),F{R0}-TODAY()<=3)'], fill=fill('FFF0C2'), font=Font(name=F, color='7A5A00', bold=True)))
ws.conditional_formatting.add(f'F{R0}:F{RN}', FormulaRule(formula=[f'AND(ISNUMBER(F{R0}),NOT({chiusa}),F{R0}<TODAY(),$A{R0}>=TODAY())'], fill=fill('F8D7DA'), font=Font(name=F, color='8A2430', bold=True)))
esempio(ws, R0, [d(6), datetime.time(10, 30), 'Sofia', 'Dott.ssa Rossi · controllo', 'Programmata', None, None, 'Esente', 'Controllo', None, None,
                 'Si sveglia tre volte a notte: è normale a questa età?', 'Mangia poco a cena', 'Quando il prossimo richiamo?', 'Da due settimane, dopo le giornate al nido', None, None, None])
esempio(ws, R0 + 1, [d(12), datetime.time(15, 0), 'Sofia', 'Allergologia · Ospedale (esempio)', 'Programmata', None, None, 'Da pagare', 'D (30/60 giorni)', d(300), 'CUP 12345678',
                     'Cosa fare se ricompare l\'orticaria?', None, None, None, None, None, None])
esempio(ws, R0 + 2, [d(-20), datetime.time(9, 0), 'Luca', 'Medico di famiglia', 'Completata', None, None, 'Esente', None, None, None, None, None, None, None, 'Tutto a posto', 'Ricontrollo tra un anno', None])

# ======================= Documenti =======================
cols = [('Documento', 26, 'in'), ('Di chi', 13, 'in'), ('Scadenza', 13, 'in'), ('Giorni', 8, 'fx'), ('Stato', 14, 'fx'),
        ('Dove si trova (carta)', 26, 'in'), ('Dove sta online', 20, 'in'), ('Note', 28, 'in')]
ws = sheet('Documenti', 'La cartellina di famiglia', 'Dove sta ogni documento e quando scade. «Da rinnovare» compare due mesi prima. Non serve scrivere i numeri dei documenti.', cols)
body(ws, cols, N); fmt(ws, 'C', N)
lista(ws, 'B', N, ref=PERSONE, strict=False)
lista(ws, 'G', N, ['Fascicolo sanitario', 'App IO', 'Mail', 'Foto sul telefono', 'Da nessuna parte'])
lista(ws, 'A', N, ['Tessera sanitaria', 'Carta d\'identità', 'Passaporto', 'Patente', 'Libretto vaccinale', 'Referto', 'Lettera di dimissione', 'Certificato', 'Altro'], strict=False)
for r in range(R0, RN + 1):
    ws[f'D{r}'] = f'=IF(C{r}="","",C{r}-TODAY())'
    ws[f'E{r}'] = f'=IF(C{r}="","",IF(D{r}<0,"Scaduto",IF(D{r}<=60,"Da rinnovare","Ok")))'
rng = f'E{R0}:E{RN}'
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'E{R0}="Scaduto"'], fill=fill('F8D7DA'), font=Font(name=F, color='8A2430', bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'E{R0}="Da rinnovare"'], fill=fill('FFF0C2'), font=Font(name=F, color='7A5A00', bold=True)))
esempio(ws, R0,     ['Tessera sanitaria', 'Sofia', d(250), None, None, 'Cartellina blu, primo scomparto', 'App IO', 'Esempio: cancella e scrivi i tuoi dati'])
esempio(ws, R0 + 1, ['Carta d\'identità', 'Luca', d(40), None, None, 'Portafoglio', 'Foto sul telefono', 'Per rinnovarla serve l\'appuntamento in Comune'])
esempio(ws, R0 + 2, ['Libretto vaccinale', 'Sofia', None, None, None, 'Cartellina blu, primo scomparto', 'Fascicolo sanitario', 'Senza scadenza: lascia vuota la data'])
for r in (R0, R0 + 1, R0 + 2): ws[f'H{r}'].font = EX

# ======================= Settimana =======================
ws = wb.create_sheet('Settimana')
ws.sheet_view.showGridLines = False
ws['A1'] = 'La settimana di tutta la famiglia'; ws['A1'].font = TITLE
ws['A2'] = 'Le cose fisse di ogni settimana e chi se ne occupa: così non le ricorda tutte una persona sola.'; ws['A2'].font = SUB
giorni = ['Lunedì', 'Martedì', 'Mercoledì', 'Giovedì', 'Venerdì', 'Sabato', 'Domenica']
fasce = ['Mattina', 'Pomeriggio', 'Sera', 'Chi se ne occupa']
ws.column_dimensions['A'].width = 16
for i, g in enumerate(giorni):
    col = L(2 + i); ws.column_dimensions[col].width = 20
    c = ws[f'{col}4']; c.value = g; c.font = H; c.fill = FILL_H; c.alignment = Alignment(horizontal='center')
for j, fa in enumerate(fasce):
    r = 5 + j; ws[f'A{r}'] = fa; ws[f'A{r}'].font = Font(name=F, bold=True, color=MALVA); ws.row_dimensions[r].height = 48
    for i in range(7):
        c = ws[f'{L(2+i)}{r}']; c.font = IN; c.border = Border(bottom=thin, right=thin); c.alignment = Alignment(wrap_text=True, vertical='top')
        if j % 2: c.fill = FILL_ALT
ws['B5'] = 'Nido 8:30'; ws['B7'] = 'Bagnetto, storia, nanna 20:30'; ws['B8'] = 'Luca mattina · Mamma sera'
ws['A10'] = 'Esempio nella colonna Lunedì: cancellalo e scrivi la tua settimana. Visite e scadenze le trovi già in «Inizia da qui».'; ws['A10'].font = EX
ws.freeze_panes = 'B5'; ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1

# ======================= Contatti =======================
cols = [('Chi', 22, 'in'), ('Nome', 24, 'in'), ('Telefono', 16, 'in'), ('Orari / quando chiamare', 28, 'in'), ('Note', 30, 'in')]
ws = sheet('Contatti', 'Contatti utili', 'Tutti i numeri che servono, in un posto solo. Stampalo e mettilo nella cartellina.', cols)
body(ws, cols, 40)
lista(ws, 'A', 40, ['Pediatra', 'Medico di famiglia', 'Guardia medica', 'CUP (prenotazioni)', 'URP della ASL', 'Fascicolo sanitario (numero verde)', 'Pronto soccorso', 'Nido / scuola', 'Nonni', 'Babysitter', 'Farmacia', 'Altro'], strict=False)
esempio(ws, R0,     ['Pediatra', 'Dott.ssa Rossi', '010 000000', 'Lun-ven 9-12, telefonate 13-14', 'Esempio: cancella e scrivi i tuoi dati'])
esempio(ws, R0 + 1, ['Guardia medica', '', '116117', 'Notte, weekend e festivi', 'Numero unico in molte regioni: verifica nella tua'])
esempio(ws, R0 + 2, ['CUP (prenotazioni)', 'CUP regionale', '', 'Per prenotare e disdire; segna anche sito e app', 'Cerca il numero della tua Regione'])
esempio(ws, R0 + 3, ['URP della ASL', '', '', 'Se la visita supera i tempi della ricetta', 'Chiedi il «percorso di tutela»'])
for r in range(R0, R0 + 4): ws[f'E{r}'].font = EX

# ======================= Prime volte =======================
cols = [('Data', 12, 'in'), ('Chi', 12, 'in'), ('Cosa è successo', 28, 'in'), ('Come è andata, con parole tue', 46, 'in'), ('Dove sta la foto o il video', 26, 'in')]
ws = sheet('Prime volte', 'Le prime volte e i ricordi', 'La prima parola, i primi passi, il primo giorno di nido. Due righe oggi valgono più di una memoria tra dieci anni.', cols)
body(ws, cols, 100, height=32); fmt(ws, 'A', 100)
lista(ws, 'B', 100, ref=PERSONE, strict=False)
esempio(ws, R0, [d(-3), 'Sofia', 'Prima parola: «pappa»', 'A cena, indicando il piatto. Abbiamo riso tutti e lei l\'ha ripetuta dieci volte.', 'Telefono di mamma, album «Sofia»'])

# ======================= Calcoli (nascosta) =======================
cs = wb.create_sheet('Calcoli')
cs.sheet_state = 'hidden'
cs['A1'] = 'chiave prossimi'; cs['B1'] = 'data'; cs['C1'] = 'cosa'; cs['D1'] = 'chi'; cs['E1'] = 'tipo'; cs['F1'] = 'chiave scadute'
cs['H1'] = 'prossimi (ordinati)'; cs['I1'] = 'scadute (ordinate)'
C0, CN = 2, 1 + 3 * N   # righe 2..901
for i in range(N):
    r = R0 + i
    # Scadenze
    c = C0 + i
    cs[f'B{c}'] = f'=IF(AND(Scadenze!E{r}<>"",Scadenze!H{r}<>"☑"),Scadenze!E{r},"")'
    cs[f'C{c}'] = f'=IF(B{c}="","",Scadenze!A{r})'
    cs[f'D{c}'] = f'=IF(B{c}="","",Scadenze!B{r})'
    cs[f'E{c}'] = f'=IF(B{c}="","","Scadenza")'
    cs[f'F{c}'] = f'=IF(B{c}="","",IF(B{c}<TODAY(),B{c}+ROW()/100000,""))'
    # Visite
    c = C0 + N + i
    cs[f'B{c}'] = f'=IF(AND(Visite!A{r}<>"",Visite!E{r}<>"Completata",Visite!E{r}<>"Annullata"),Visite!A{r},"")'
    cs[f'C{c}'] = (f'=IF(B{c}="","","Visita"&IF(Visite!D{r}="",""," · "&Visite!D{r})'
                   f'&IF(Visite!B{r}="","",IF(ISNUMBER(Visite!B{r})," · ore "&TEXT(Visite!B{r},"hh:mm")," · "&Visite!B{r})))')
    cs[f'D{c}'] = f'=IF(B{c}="","",Visite!C{r})'
    cs[f'E{c}'] = f'=IF(B{c}="","","Visita")'
    cs[f'F{c}'] = '=""'
    # Documenti (entro 90 giorni o scaduti)
    c = C0 + 2 * N + i
    cs[f'B{c}'] = f'=IF(Documenti!C{r}="","",IF(Documenti!D{r}<=90,Documenti!C{r},""))'
    cs[f'C{c}'] = f'=IF(B{c}="","","Rinnovo · "&Documenti!A{r})'
    cs[f'D{c}'] = f'=IF(B{c}="","",Documenti!B{r})'
    cs[f'E{c}'] = f'=IF(B{c}="","","Documento")'
    cs[f'F{c}'] = f'=IF(B{c}="","",IF(B{c}<TODAY(),B{c}+ROW()/100000,""))'
for c in range(C0, CN + 1):
    cs[f'A{c}'] = f'=IF(B{c}="","",IF(B{c}>=TODAY(),B{c}+ROW()/100000,""))'
KEYS, OVER = f'$A${C0}:$A${CN}', f'$F${C0}:$F${CN}'
NP, NS = 10, 5
for k in range(1, NP + 1):
    cs[f'H{1+k}'] = f'=IF(COUNT({KEYS})>={k},SMALL({KEYS},{k}),"")'
for k in range(1, NS + 1):
    cs[f'I{1+k}'] = f'=IF(COUNT({OVER})>={k},LARGE({OVER},{k}),"")'

# ======================= Inizia da qui =======================
ws = wb.active; ws.title = 'Inizia da qui'
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDE', [13, 40, 14, 16, 4]): ws.column_dimensions[col].width = w
ws['A1'] = 'Il gestionale di famiglia'; ws['A1'].font = Font(name='Georgia', bold=True, size=22, color=INK); ws.row_dimensions[1].height = 34
ws['A2'] = f'di serenainmente.it · versione {VERSIONE} · i tuoi dati restano sul tuo telefono, computer o Drive'; ws['A2'].font = SUB
ws['A4'] = '=IF(HOUR(NOW())<13,"Buongiorno!",IF(HOUR(NOW())<18,"Buon pomeriggio!","Buonasera!"))'
ws['A4'].font = Font(name='Georgia', size=16, color=MALVA); ws.row_dimensions[4].height = 26
ws['A5'] = 'Oggi'; ws['A5'].font = Font(name=F, color=MUTED)
ws['B5'] = '=TODAY()'; ws['B5'].number_format = 'DD/MM/YYYY'; ws['B5'].font = Font(name=F, bold=True, color=INK); ws['B5'].alignment = Alignment(horizontal='left')

def titolo(r, t):
    ws[f'A{r}'] = t; ws[f'A{r}'].font = Font(name='Georgia', bold=True, size=13, color=INK); ws.row_dimensions[r].height = 24

titolo(7, 'In breve')
SC, VI, DO = 'Scadenze', 'Visite', 'Documenti'
cards = [
    (f'=COUNTIFS({SC}!$F${R0}:$F${RN},">=0",{SC}!$F${R0}:$F${RN},"<=30")', 'scadenze nei prossimi 30 giorni', ROSA, MALVA),
    (f'=COUNTIFS(Calcoli!$E${C0}:$E${CN},"Visita",Calcoli!$B${C0}:$B${CN},">="&TODAY(),Calcoli!$B${C0}:$B${CN},"<="&(TODAY()+14))', 'visite nelle prossime due settimane', SALVIA, SALVIA_T),
    (f'=COUNTIF({DO}!$E${R0}:$E${RN},"Scaduto")+COUNTIF({DO}!$E${R0}:$E${RN},"Da rinnovare")', 'documenti da rinnovare (scaduti o entro due mesi)', SABBIA, SABBIA_T),
    (f'=COUNTIF({SC}!$G${R0}:$G${RN},"Scaduta")', 'scadenze passate e non ancora fatte', 'FFFFFF', INK),
    (f'=COUNTIFS({VI}!$H${R0}:$H${RN},"Da pagare",{VI}!$E${R0}:$E${RN},"<>Annullata")', 'ticket ancora da pagare', 'FFFFFF', INK),
]
for i, (f, lab, bg, fg) in enumerate(cards):
    r = 8 + i
    ws[f'A{r}'] = f; ws[f'A{r}'].font = Font(name='Georgia', bold=True, size=20, color=fg)
    ws[f'A{r}'].alignment = Alignment(horizontal='center', vertical='center')
    ws[f'B{r}'] = lab; ws[f'B{r}'].font = Font(name=F, size=11, color=INK); ws[f'B{r}'].alignment = Alignment(vertical='center', wrap_text=True)
    ws.merge_cells(f'B{r}:D{r}')
    for col in 'ABCD':
        ws[f'{col}{r}'].fill = fill(bg); ws[f'{col}{r}'].border = Border(bottom=Side(style='thin', color='FFFFFF' if bg != 'FFFFFF' else LINE))
    ws.row_dimensions[r].height = 34
ws.conditional_formatting.add('A11', FormulaRule(formula=['A11>0'], fill=fill('F8D7DA'), font=Font(name='Georgia', bold=True, size=20, color='8A2430')))

titolo(14, 'Prossimi impegni')
ws['A15'] = 'Visite, scadenze e rinnovi, dal più vicino. Si aggiorna da solo.'; ws['A15'].font = SUB
for i, h in enumerate(['Quando', 'Cosa', 'Per chi', 'Tra']):
    c = ws.cell(row=16, column=1 + i, value=h); c.font = Font(name=F, bold=True, size=9, color=MUTED); c.border = Border(bottom=thin)
def riga_lista(r, keycell, empty_msg=None, scadute=False):
    look = lambda col: f'INDEX(Calcoli!${col}${C0}:${col}${CN},MATCH({keycell},Calcoli!${"F" if scadute else "A"}${C0}:${"F" if scadute else "A"}${CN},0))'
    ws[f'A{r}'] = f'=IF({keycell}="","",{look("B")})'
    msg = f'"{empty_msg}"' if empty_msg else '""'
    ws[f'B{r}'] = f'=IF({keycell}="",{msg},{look("C")})'
    ws[f'C{r}'] = f'=IF({keycell}="","",{look("D")})'
    if scadute:
        ws[f'D{r}'] = f'=IF(A{r}="","","da "&(TODAY()-A{r})&IF(TODAY()-A{r}=1," giorno"," giorni"))'
    else:
        ws[f'D{r}'] = f'=IF(A{r}="","",IF(A{r}=TODAY(),"oggi",IF(A{r}=TODAY()+1,"domani","tra "&(A{r}-TODAY())&" giorni")))'
    ws[f'A{r}'].number_format = 'DD/MM'
    ws[f'A{r}'].font = Font(name=F, bold=True, size=11, color=MALVA if not scadute else '8A2430')
    ws[f'B{r}'].font = Font(name=F, size=11, color=INK); ws[f'C{r}'].font = Font(name=F, size=10, color=MUTED)
    ws[f'D{r}'].font = Font(name=F, size=10, color=MUTED)
    for col in 'ABCD':
        ws[f'{col}{r}'].border = Border(bottom=thin); ws[f'{col}{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=(col == 'B'))
    ws.row_dimensions[r].height = 24
for k in range(1, NP + 1):
    riga_lista(16 + k, f'Calcoli!$H${1+k}', 'Nessun impegno in arrivo: aggiungi una visita o una scadenza nelle loro schede.' if k == 1 else None)
ws.conditional_formatting.add(f'A17:D{16+NP}', FormulaRule(formula=['AND(ISNUMBER($A17),$A17-TODAY()<=2)'], fill=fill('FFF6DB')))

r0 = 18 + NP
titolo(r0, 'Da controllare')
ws[f'A{r0+1}'] = 'Scadenze passate senza «Fatto» e documenti già scaduti, dalla più recente.'; ws[f'A{r0+1}'].font = SUB
for k in range(1, NS + 1):
    riga_lista(r0 + 1 + k, f'Calcoli!$I${1+k}', 'Niente in sospeso.' if k == 1 else None, scadute=True)

r1 = r0 + NS + 3
titolo(r1, 'Come si usa')
guida = [
    ('1. Famiglia', 'Scrivi i nomi di casa una volta sola: compaiono nei menu di tutte le schede.'),
    ('2. Scadenze', 'Tutto ciò che ha una data. Se si ripete, scegli «Ogni mese» o «Ogni anno»: la data si sposta da sola.'),
    ('3. Visite', 'Data, per chi, dove. «Disdire entro» si calcola; segna «Completata» o «Annullata» e la visita esce dagli impegni.'),
    ('4. Documenti', 'Dove sta ogni documento e quando scade: due mesi prima compare qui sopra.'),
    ('E poi', 'Settimana, Contatti e Prime volte: da compilare quando vuoi.'),
]
for i, (k, v) in enumerate(guida):
    r = r1 + 1 + i
    ws[f'A{r}'] = k; ws[f'A{r}'].font = Font(name=F, bold=True, color=MALVA); ws[f'A{r}'].alignment = Alignment(vertical='top')
    ws[f'B{r}'] = v; ws[f'B{r}'].font = Font(name=F, color=INK); ws[f'B{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(f'B{r}:D{r}'); ws.row_dimensions[r].height = 30
r = r1 + len(guida) + 2
ws[f'A{r}'] = 'Legenda'; ws[f'A{r}'].font = Font(name=F, bold=True, color=INK)
ws[f'A{r+1}'] = 'Testo blu'; ws[f'A{r+1}'].font = IN; ws[f'B{r+1}'] = 'Celle da compilare (le righe di esempio vanno cancellate)'
ws[f'A{r+2}'] = 'Testo nero'; ws[f'A{r+2}'].font = FX; ws[f'B{r+2}'] = 'Si calcola da solo: non scriverci sopra'
for x in (r + 1, r + 2): ws[f'B{x}'].font = Font(name=F, color=INK)
ws[f'A{r+4}'] = ('Il file non manda notifiche: ti mostra cosa arriva ogni volta che lo apri. Resta dove lo salvi tu (telefono, computer, OneDrive o Google Drive) '
                 'e nessun dato arriva a Serena. Non è un archivio medico: scrivi solo quello che serve. Per vaccini e terapie il riferimento resta il pediatra o il medico.')
ws[f'A{r+4}'].font = Font(name=F, size=9, color=MUTED); ws[f'A{r+4}'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells(f'A{r+4}:D{r+4}'); ws.row_dimensions[r + 4].height = 54
ws[f'A{r+5}'] = 'serenainmente.it/guide'; ws[f'A{r+5}'].font = Font(name=F, size=9, color=MALVA, underline='single'); ws[f'A{r+5}'].hyperlink = 'https://serenainmente.it/guide/'
ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = 0
wb.properties.title = 'Il gestionale di famiglia · serenainmente.it'
wb.properties.creator = 'Serena'
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT)
